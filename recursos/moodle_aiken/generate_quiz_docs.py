import docx
import fitz  # PyMuPDF
import re

aiken_path = r"c:\Users\crist\Downloads\Docker\GUIAS_11VA_EDICION\moodle_aiken\Quiz_Sesiones_1_a_4_Aiken.txt"
docx_path = r"c:\Users\crist\Downloads\Docker\GUIAS_11VA_EDICION\moodle_aiken\Quiz_Sesiones_1_a_4.docx"
pdf_path = r"c:\Users\crist\Downloads\Docker\GUIAS_11VA_EDICION\moodle_aiken\Quiz_Sesiones_1_a_4.pdf"
txt_clean_path = r"c:\Users\crist\Downloads\Docker\GUIAS_11VA_EDICION\moodle_aiken\Quiz_Sesiones_1_a_4_Texto.txt"

with open(aiken_path, "r", encoding="utf-8") as f:
    lines = [line.strip() for line in f.readlines()]

questions = []
cur_q = None

for line in lines:
    if not line:
        continue
    if line.startswith("ANSWER:"):
        ans_letter = line.replace("ANSWER:", "").strip().upper()
        if cur_q:
            cur_q["answer"] = ans_letter
            questions.append(cur_q)
            cur_q = None
    elif re.match(r"^[A-E]\.\s+", line):
        if cur_q:
            cur_q["options"].append(line)
    else:
        if cur_q is None:
            cur_q = {"question": line, "options": [], "answer": ""}
        else:
            cur_q["question"] += " " + line

# 1. Create Word Document (.docx)
doc = docx.Document()
doc.add_heading("Cuestionario de Docker - Sesiones 1 a 4", level=1)

for i, q in enumerate(questions, 1):
    doc.add_heading(f"{i}. {q['question']}", level=3)
    for opt in q["options"]:
        doc.add_paragraph(opt)
    doc.add_paragraph(f"Respuesta correcta: {q['answer']}")
    doc.add_paragraph("")  # empty space

doc.save(docx_path)
print(f"Created Word: {docx_path}")

# 2. Create PDF with PyMuPDF (fitz)
pdf_doc = fitz.open()
page = pdf_doc.new_page(width=595, height=842) # A4
font_size = 10
line_height = 14
y = 50
margin = 40

def check_page_and_draw(text, is_heading=False, is_bold=False):
    global y, page, pdf_doc
    if y > 780:
        page = pdf_doc.new_page(width=595, height=842)
        y = 50
    size = 12 if is_heading else 9.5
    # PyMuPDF insert_textbox with wrap
    rect = fitz.Rect(margin, y, 595 - margin, y + 100)
    res = page.insert_textbox(rect, text, fontsize=size, fontname="helv" if not is_bold else "hebo")
    # approximate height
    lines_count = max(1, len(text) // 80 + 1)
    y += lines_count * line_height + (6 if is_heading else 2)

check_page_and_draw("Cuestionario de Docker - Sesiones 1 a 4", is_heading=True, is_bold=True)
y += 10

for i, q in enumerate(questions, 1):
    check_page_and_draw(f"{i}. {q['question']}", is_heading=True, is_bold=True)
    for opt in q["options"]:
        check_page_and_draw(f"   {opt}")
    check_page_and_draw(f"   Respuesta correcta: {q['answer']}", is_bold=True)
    y += 8

pdf_doc.save(pdf_path)
pdf_doc.close()
print(f"Created PDF: {pdf_path}")

# 3. Create Clean Text (.txt)
with open(txt_clean_path, "w", encoding="utf-8") as f:
    f.write("Cuestionario de Docker - Sesiones 1 a 4\n\n")
    for i, q in enumerate(questions, 1):
        f.write(f"{i}. {q['question']}\n")
        for opt in q["options"]:
            f.write(f"{opt}\n")
        f.write(f"Respuesta correcta: {q['answer']}\n\n")

print(f"Created Text: {txt_clean_path}")
